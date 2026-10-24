from datetime import date, timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from apps.companies.models import Company, InterviewQuestion, Role, Round

from . import ats, discovery, feeds, skillgraph
from .ats import ats_check
from .models import CandidateProfile, Job, JobApplication, JobBoard, JobFeedState, ResumeCheck, SavedJob
from .practice import score_answer

User = get_user_model()
RESUME = "Software engineer with 4 years building web apps using React, Python and Django. " * 3


def make_job(**kw):
    defaults = dict(source="t", external_id=str(Job.objects.count()), title="Python Developer",
                    company="Acme", url="https://x.test/1", posted_at=timezone.now(),
                    description="We need Python, Django, PostgreSQL, Docker and AWS experience.")
    defaults.update(kw)
    return Job.objects.create(**defaults)


class SkillGraphATSTests(TestCase):
    def test_alias_and_implication(self):
        r = ats_check("I know k8s and reactjs", "Frontend", "Need Docker, JavaScript and React")
        self.assertIn("javascript", r["matched_keywords"])  # implied by React
        self.assertIn("docker", r["matched_keywords"])      # implied by Kubernetes alias
        self.assertEqual(r["implied_keywords"], ["docker", "javascript"])
        keywords = next(b for b in r["breakdown"] if b["key"] == "keywords")["score"]
        self.assertTrue(50 < keywords < 100)  # implied skills earn partial credit
        self.assertTrue(any("JavaScript" in t for t in r["suggestions"]))

    def test_missing_keywords_and_score(self):
        r = ats_check(RESUME, "Python Developer", "Python Django PostgreSQL Docker AWS")
        self.assertEqual(set(r["missing_keywords"]), {"docker", "aws", "postgresql"})
        self.assertLess(r["ats_score"], 100)

    def test_no_recognisable_skills(self):
        self.assertIsNone(ats_check(RESUME, "Role", "Be nice")["ats_score"])

    def test_word_boundaries(self):
        r = ats_check("javascript developer", "x", "Need Java")
        self.assertEqual(r["matched_keywords"], [])

    def test_plain_english_words_are_not_skills(self):
        text = "Ready to go-to-market. Excel at communication. React to incidents. Send your CV by 5 pm. R&D team."
        self.assertEqual(skillgraph.extract_skills(text), {"communication"})
        self.assertEqual(skillgraph.extract_skills("Go, REST, MS Excel, React, Spring Boot"),
                         {"go", "rest api", "excel", "react", "spring boot"})


JD = """About Acme
We are a fintech company. Our sales and marketing teams live in Salesforce.
What you'll do
- Build backend services in Go and Python
Requirements
- 3+ years of experience in backend development
- Strong proficiency in Python or Java, PostgreSQL and Redis
- Experience with Docker and AWS
- Bachelor's degree in Computer Science or equivalent
Nice to have
- Terraform, GraphQL
Benefits
- Health insurance and Excel training
"""

CV = """Priya Sharma
priya@example.com | +91 98765 43210
Summary
Backend engineer building payment APIs.
Experience
Software Engineer, Acme Corp
Jan 2023 - Present
- Built REST APIs with Python, Django and PostgreSQL serving 2M requests/day, cut latency by 40%
- Deployed on AWS with Docker; 10k events/sec through Redis
Education
B.Tech Computer Science, IIT Delhi
2018 - 2022
Skills
Python, Django, PostgreSQL, Docker, Redis, AWS
"""


class ATSScoringTests(TestCase):
    TODAY = date(2026, 10, 1)

    def check(self, cv=CV, title="Backend Engineer", jd=JD):
        return ats_check(cv, title, jd, today=self.TODAY)

    def kw(self, r):
        return {k["skill"]: (k["importance"], k["status"]) for k in r["keywords"]}

    def test_sections_set_importance(self):
        kw = self.kw(self.check())
        self.assertEqual(kw["redis"], ("required", "matched"))
        self.assertEqual(kw["go"], ("general", "missing"))
        self.assertEqual(kw["terraform"], ("preferred", "missing"))
        for skill in ("salesforce", "sales", "marketing", "excel"):  # About us / Benefits only
            self.assertNotIn(skill, kw)

    def test_either_or_requirement(self):
        kw = self.kw(self.check())
        self.assertEqual(kw["python"], ("required", "matched"))
        self.assertNotIn("java", kw)
        self.assertNotIn("java", self.check()["missing_keywords"])
        r = self.check(cv=CV.replace("Python, ", ""))  # Django still implies Python
        self.assertEqual(self.kw(r)["python"], ("required", "implied"))
        r = self.check(cv=CV.replace("Python, ", "").replace("Django", "Rails"))
        self.assertIn("java or python", r["missing_keywords"])

    def test_missing_required_costs_more_than_nice_to_have(self):
        base = self.check()["ats_score"]
        no_redis = self.check(cv=CV.replace("Redis", "Memcached"))["ats_score"]
        with_terraform = self.check(cv=CV + "Terraform\n")["ats_score"]
        self.assertLess(no_redis, base - 3)
        self.assertLessEqual(with_terraform - base, 3)

    def test_experience_from_dates_excludes_education(self):
        self.assertEqual(ats.required_years(JD), 3)
        self.assertEqual(ats.required_years("Experience: 2-4 years in Python"), 2)
        self.assertIsNone(ats.required_years("Founded 25 years ago. Python."))
        self.assertEqual(ats.resume_years(CV, self.TODAY), 3.8)  # Jan 2023 -> Oct 2026, not 2018-2022
        self.assertEqual(ats.resume_years("Summary\n5+ years of professional experience in Java"), 5)
        r = self.check(cv=CV.replace("Jan 2023", "Jan 2025"))
        self.assertEqual(r["experience"], {"required_years": 3, "resume_years": 1.8})
        self.assertLess(next(b["score"] for b in r["breakdown"] if b["key"] == "experience"), 100)
        self.assertTrue(any("3+ years" in t for t in r["suggestions"]))

    def test_degree_requirement(self):
        self.assertEqual(ats.required_degree(JD), (1, True))
        self.assertEqual(ats.required_degree("You must be in the Bangalore office."), (None, False))
        self.assertEqual(ats.required_degree("MS in CS or PhD"), (2, False))
        r = self.check(cv=CV.replace("B.Tech Computer Science, IIT Delhi", "Delhi Public School"))
        self.assertLess(next(b["score"] for b in r["breakdown"] if b["key"] == "education"), 100)

    def test_title_match(self):
        self.assertEqual(ats._core_title("Senior Backend Developer - Payments (Remote)"), ["backend", "engineer"])
        self.assertEqual(ats._core_title("SDE-2"), ["engineer"])
        r = self.check(title="Data Scientist")
        self.assertLess(next(b["score"] for b in r["breakdown"] if b["key"] == "title"), 100)
        self.assertEqual(next(b["score"] for b in self.check()["breakdown"] if b["key"] == "title"), 100)

    def test_structure_checks(self):
        r = self.check(cv="python django postgresql redis docker aws " * 10)
        structure = next(b["score"] for b in r["breakdown"] if b["key"] == "structure")
        self.assertLess(structure, 30)
        tips = " ".join(r["suggestions"])
        self.assertIn("email and phone", tips)
        self.assertIn("section headings", tips)

    def test_company_name_and_team_mentions_are_not_requirements(self):
        jd = ("At Databricks we build the lakehouse. You'll work closely with product management and sales.\n"
              "Requirements\n- Python and SQL\n- Spark\n")
        kw = self.kw(ats_check(CV, "Software Engineer", jd, company="Databricks", today=self.TODAY))
        self.assertEqual(set(kw), {"python", "sql", "spark"})

    def test_real_world_headings(self):
        for line, kind in [("What you’ll bring", "required"), ("What skills do I need?", "required"),
                           ("Bonus if you have the following", "preferred"), ("A typical day:", "general"),
                           ("Our commitment to diversity and inclusion", "ignore"), ("About Groww:", "ignore"),
                           ("About you", "required"), ("Experience with Python", None),
                           ("Strong communication skills", None)]:
            self.assertEqual((ats._heading(line) or [None])[0], kind, line)
        lines = list(ats._posting_lines("Requirements\n- Location: Bangalore\n- Python and Go\n"))
        self.assertEqual(lines[-1], ("required", "- Python and Go"))  # a one-off label doesn't open a section

    def test_parts_without_requirements_are_left_out(self):
        r = ats_check(CV, "Backend Engineer", "Python, Docker and AWS")
        self.assertEqual({b["key"] for b in r["breakdown"]}, {"keywords", "title", "structure"})
        self.assertEqual(sum(b["weight"] for b in r["breakdown"]), 100)


class PracticeScoringTests(TestCase):
    IDEAL = "Use a hash map to store seen values and look up the complement in constant time"

    def test_blank_is_zero(self):
        self.assertEqual(score_answer("", self.IDEAL)["score"], 0)

    def test_good_answer_scores_high(self):
        ans = ("I would iterate once and store each seen value in a hash map, then look up the "
               "complement of the current number; the lookup is constant time so overall it is linear.")
        self.assertGreaterEqual(score_answer(ans, self.IDEAL)["score"], 8)

    def test_short_answer_capped(self):
        self.assertLessEqual(score_answer("hash map seen values complement constant time lookup", self.IDEAL)["score"], 4)


@override_settings(JOB_REFRESH_SYNC=True)
@mock.patch.object(feeds, "refresh_boards", lambda: (0, 0))
class FeedRefreshTests(TestCase):
    def test_lazy_refresh_every_five_hours(self):
        row = feeds._row(source="remotive", external_id="1", title="T", company="C", location="", url="https://a.test",
                         description="", tags=[], posted_at=timezone.now())
        with mock.patch.object(feeds, "FETCHERS", (lambda: [dict(row)],)):
            self.assertTrue(feeds.refresh_if_stale())
            self.assertFalse(feeds.refresh_if_stale())  # fresh -> no refetch
            JobFeedState.objects.update(last_refreshed_at=timezone.now() - timedelta(hours=5, minutes=1))
            self.assertTrue(feeds.refresh_if_stale())
        self.assertEqual(Job.objects.count(), 1)  # upsert, not duplicate
        self.assertFalse(feeds.is_sweeping(JobFeedState.objects.get()))

    def test_failing_feed_does_not_raise(self):
        def boom():
            raise RuntimeError("down")
        with mock.patch.object(feeds, "FETCHERS", (boom,)):
            self.assertTrue(feeds.refresh_if_stale())
        self.assertIn("boom", JobFeedState.objects.get().last_error)

    def test_background_sweep_when_not_sync(self):
        with override_settings(JOB_REFRESH_SYNC=False), mock.patch.object(feeds.threading, "Thread") as thread:
            self.assertTrue(feeds.refresh_if_stale())
        thread.assert_called_once()
        thread.return_value.start.assert_called_once()
        self.assertTrue(feeds.is_sweeping(JobFeedState.objects.get()))


    def test_external_sweeper_requests_only_cover_when_twice_stale(self):
        JobFeedState.objects.create(pk=1, last_refreshed_at=timezone.now() - timedelta(hours=6),
                                    last_finished_at=timezone.now() - timedelta(hours=6))
        with override_settings(JOB_SWEEPER_EXTERNAL=True), mock.patch.object(feeds, "FETCHERS", ()):
            self.assertFalse(feeds.refresh_if_stale())  # the sweeper's job, not the request's
            JobFeedState.objects.update(last_refreshed_at=timezone.now() - timedelta(hours=11))
            self.assertTrue(feeds.refresh_if_stale())

    def test_sweep_that_died_midway_is_reclaimed(self):
        JobFeedState.objects.create(pk=1, last_refreshed_at=timezone.now() - timedelta(hours=2),
                                    last_finished_at=timezone.now() - timedelta(hours=7))
        self.assertTrue(feeds.claim_sweep(timedelta(hours=5)))
        self.assertFalse(feeds.claim_sweep(timedelta(hours=5)))  # just claimed, still running

    def test_sweep_jobs_command_once(self):
        from django.core.management import call_command
        with mock.patch.object(feeds, "FETCHERS", ()):
            call_command("sweep_jobs", "--once", stdout=mock.Mock())
            state = JobFeedState.objects.get()
            self.assertFalse(feeds.is_sweeping(state))
            call_command("sweep_jobs", "--once", stdout=mock.Mock())  # fresh -> no second sweep
        self.assertEqual(JobFeedState.objects.get().last_refreshed_at, state.last_refreshed_at)

    def test_resweep_skips_unchanged_rows(self):
        row = feeds._row(source="remotive", external_id="1", title="T", company="C", location="", url="https://a.test",
                         description="Python", tags=[], posted_at=timezone.now())
        feeds._upsert([dict(row)])
        with mock.patch.object(Job.objects, "bulk_create") as bulk:
            feeds._upsert([dict(row)])
        bulk.assert_not_called()
        feeds._upsert([dict(row, description="Python and Go")])
        self.assertEqual(Job.objects.get().description, "Python and Go")
        self.assertEqual(Job.objects.count(), 1)


def _resp(payload):
    r = mock.Mock()
    r.json.return_value = payload
    r.raise_for_status.return_value = r
    return r


class DescriptionCleaningTests(TestCase):
    def test_html_keeps_paragraphs_and_bullets(self):
        raw = "<h3>About us</h3><p>We build &amp; ship.</p><ul><li><p>Python</p></li><li>Docker</li></ul>"
        self.assertEqual(feeds._clean_description(raw), "About us\nWe build & ship.\n- Python\n- Docker")

    def test_plain_text_newlines_survive(self):
        raw = feeds._html_lines("Line one\n\nUse <Go>\n- x")
        self.assertEqual(feeds._clean_description(raw), "Line one\nUse <Go>\n- x")


class CompanyBoardTests(TestCase):
    """Greenhouse / Lever / Ashby boards: parsing, closed-role removal, auto-deactivation."""

    GH = {"jobs": [
        {"id": 1, "title": "Senior Backend Engineer", "absolute_url": "https://gh.test/1",
         "location": {"name": "Bengaluru, India"}, "first_published": "2026-10-20T10:00:00Z", "company_name": "Acme"},
        {"id": 2, "title": "Software Engineering Intern", "absolute_url": "https://gh.test/2",
         "location": {"name": "Remote - India"}, "updated_at": "2026-10-21T10:00:00Z"},
    ]}
    LEVER = [{"id": "abc", "text": "Data Analyst", "hostedUrl": "https://lever.test/abc", "createdAt": 1792000000000,
              "workplaceType": "hybrid", "categories": {"location": "Pune", "team": "Data", "commitment": "Full-time"},
              "descriptionPlain": "Need SQL, Python and Tableau. 2+ years of experience.", "lists": []}]
    ASHBY = {"jobs": [{"id": "z1", "title": "Frontend Engineer", "jobUrl": "https://ashby.test/z1", "isListed": True,
                       "location": "Remote", "isRemote": True, "workplaceType": "Remote",
                       "publishedAt": "2026-10-19T00:00:00Z", "descriptionPlain": "React and TypeScript."}]}

    SR = {"totalFound": 2, "content": [
        {"id": "601", "name": "Delivery Operations Manager", "company": {"identifier": "SWIGGY"},
         "releasedDate": "2026-10-22T06:00:00Z", "location": {"city": "Bengaluru", "country": "in", "hybrid": True},
         "function": {"label": "Operations"}, "typeOfEmployment": {"label": "Full-time"}},
        {"id": "602", "name": "Legacy Role", "company": {"identifier": "SWIGGY"},
         "releasedDate": "2026-01-01T00:00:00Z", "location": {"city": "Pune", "country": "in"}},
    ]}

    def setUp(self):
        self.gh = JobBoard.objects.create(ats="greenhouse", token="acme", company="Acme")
        self.lv = JobBoard.objects.create(ats="lever", token="beta", company="Beta")
        self.ab = JobBoard.objects.create(ats="ashby", token="gamma", company="Gamma")
        self.sr = JobBoard.objects.create(ats="smartrecruiters", token="Swiggy", company="Swiggy")

    def _client(self, gh=None):
        client = mock.Mock()
        payloads = {"greenhouse": gh or self.GH, "lever": self.LEVER, "ashby": self.ASHBY}
        client.get.side_effect = lambda url, **kw: _resp(next(v for k, v in (
            ("boards-api.greenhouse.io", payloads["greenhouse"]), ("api.lever.co", payloads["lever"]),
            ("api.ashbyhq.com", payloads["ashby"]), ("api.smartrecruiters.com", self.SR)) if k in url))
        return client

    def _sweep(self, client):
        with mock.patch.object(feeds.httpx, "Client") as cls:
            cls.return_value.__enter__.return_value = client
            return feeds.refresh_boards(list(JobBoard.objects.filter(is_active=True)))

    def test_parses_all_ats(self):
        self.assertEqual(self._sweep(self._client()), (4, 0))
        senior = Job.objects.get(external_id="acme:1")
        self.assertEqual((senior.level, senior.work_mode, senior.board, senior.company),
                         ("senior", "onsite", self.gh, "Acme"))
        intern = Job.objects.get(external_id="acme:2")
        self.assertEqual((intern.level, intern.work_mode), ("entry", "remote"))
        analyst = Job.objects.get(source="lever")
        self.assertEqual((analyst.work_mode, analyst.level), ("hybrid", "mid"))
        self.assertIn("tableau", analyst.skills)
        fe = Job.objects.get(source="ashby")
        self.assertEqual(fe.work_mode, "remote")
        self.assertIn("react", fe.skills)

    def test_closed_roles_removed_but_tracking_kept(self):
        self._sweep(self._client())
        user = User.objects.create_user(username="s", email="s@ex.com", password="pw12345!")
        job = Job.objects.get(external_id="acme:2")
        SavedJob.objects.create(user=user, job=job, **SavedJob.snapshot_of(job))
        self._sweep(self._client(gh={"jobs": self.GH["jobs"][:1]}))  # company closed job 2
        self.assertFalse(Job.objects.filter(external_id="acme:2").exists())
        saved = SavedJob.objects.get()
        self.assertIsNone(saved.job)
        self.assertEqual(saved.title, "Software Engineering Intern")

    def test_failing_board_deactivates_after_repeated_failures(self):
        client = self._client()
        client.get.side_effect = RuntimeError("404")
        for _ in range(feeds.DEACTIVATE_AFTER):
            self._sweep(client)
        self.assertFalse(JobBoard.objects.filter(is_active=True).exists())

    def test_smartrecruiters_india_roles_and_lazy_description(self):
        self._sweep(self._client())
        job = Job.objects.get(source="smartrecruiters")  # the January posting is past JOB_MAX_AGE_DAYS
        self.assertEqual((job.location, job.work_mode, job.in_india, job.company),
                         ("Bengaluru, India", "hybrid", True, "Swiggy"))
        self.assertEqual(job.url, "https://jobs.smartrecruiters.com/SWIGGY/601")
        detail = {"jobAd": {"sections": {"jobDescription": {"title": "Job Description",
                                                            "text": "<ul><li>SQL and Excel</li></ul>"}}}}
        with mock.patch.object(feeds.httpx, "get", return_value=_resp(detail)) as get:
            feeds.ensure_description(job)
        self.assertIn("/companies/Swiggy/postings/601", get.call_args.args[0])
        job.refresh_from_db()
        self.assertEqual(job.description, "Job Description\n- SQL and Excel")
        self.assertIn("sql", job.skills)

    def test_india_flag(self):
        self._sweep(self._client())
        flags = dict(Job.objects.values_list("external_id", "in_india"))
        self.assertTrue(flags["acme:1"])        # Bengaluru, India
        self.assertTrue(flags["acme:2"])        # Remote - India
        self.assertTrue(flags["beta:abc"])      # Pune
        self.assertFalse(flags["gamma:z1"])     # plain "Remote" (could be US-only)
        self.assertTrue(feeds.is_india_job("Remote - APAC", "remote"))
        self.assertFalse(feeds.is_india_job("APAC", "onsite"))

    def test_old_postings_skipped_and_deleted(self):
        stale = make_job(posted_at=timezone.now() - timedelta(days=61), board=self.ab, source="ashby",
                         external_id="gamma:old")
        aggregator = make_job(posted_at=timezone.now() - timedelta(days=61))
        fresh = make_job(posted_at=timezone.now() - timedelta(days=59))
        with mock.patch.object(feeds, "refresh_jobs", lambda: (0, [])), \
                mock.patch.object(feeds, "refresh_boards", lambda: (0, 0)):
            feeds.run_sweep()
        self.assertFalse(Job.objects.filter(pk__in=[stale.pk, aggregator.pk]).exists())
        self.assertTrue(Job.objects.filter(pk=fresh.pk).exists())

    def test_greenhouse_description_fetched_lazily(self):
        self._sweep(self._client())
        job = Job.objects.get(external_id="acme:1")
        self.assertEqual(job.description, "")
        with mock.patch.object(feeds.httpx, "get", return_value=_resp(
                {"content": "&lt;p&gt;Go, Kubernetes and PostgreSQL. 6+ years.&lt;/p&gt;"})):
            feeds.ensure_description(job)
        job.refresh_from_db()
        self.assertIn("Kubernetes", job.description)
        self.assertIn("kubernetes", job.skills)
        self._sweep(self._client())  # re-sweep must not wipe the fetched description
        job.refresh_from_db()
        self.assertIn("kubernetes", job.skills)

    def test_boards_file_seeds_once(self):
        self.assertGreaterEqual(feeds.sync_boards_from_file(), 900)
        self.assertEqual(feeds.sync_boards_from_file(), 0)


class IndiaBoardTests(TestCase):
    """Keka / Workday / Freshteam boards and the Indian-company board discovery."""

    KEKA_PAGE = "<script>fetch('/ats/documents/b5279857-cf81-4dde-a215-fc48957ee2b5/careerportal/x.html')</script>"
    KEKA = [{"id": 145340, "title": "Security Engineer 3", "description": "<div>AWS and Python. 5+ years.</div>",
             "departmentName": "Engineering", "experience": "5-6 Years", "publishedOn": "2026-10-22T10:19:10.517Z",
             "jobLocations": [{"city": "Bengaluru", "countryCode": "IN"}, {"city": "Pune", "countryCode": "IN"}]}]
    FRESHTEAM = {"jobs": [{"id": 6000168140, "title": "Backend Engineer", "description": "<p>Go and Kafka</p>",
                           "deleted": False, "created_at": "2026-10-20T08:38:00.000Z", "remote": False,
                           "url": "https://haptik.freshteam.com/jobs/297x/backend", "branch_id": 7, "job_role_id": 9},
                          {"id": 1, "title": "Sample Job Posting", "deleted": False, "created_at": "2026-10-20T08:38:00Z",
                           "url": "https://haptik.freshteam.com/jobs/1/sample", "branch_id": 7}],
                 "branches": [{"id": 7, "city": "Mumbai", "country_code": "IN"}],
                 "job_roles": [{"id": 9, "name": "Engineering"}]}
    WD_FACETS = [{"facetParameter": "locationMainGroup", "values": [
        {"facetParameter": "locationHierarchy1", "values": [
            {"descriptor": "United States", "id": "us1", "count": 9}, {"descriptor": "India", "id": "in1", "count": 2}]}]}]
    WD_JOBS = [
        {"title": "SoC Architect", "externalPath": "/job/India-Bengaluru/SoC-Architect_JR1", "locationsText": "India, Bengaluru",
         "postedOn": "Posted Today"},
        {"title": "Research Scientist", "externalPath": "/job/India-Pune/Research-Scientist_JR2",
         "locationsText": "3 Locations", "postedOn": "Posted 30+ Days Ago"},
    ]

    def _client(self, keka_page=KEKA_PAGE):
        client = mock.Mock()
        page = mock.Mock(text=keka_page)
        page.raise_for_status.return_value = page
        client.get.side_effect = lambda url, **kw: (
            page if url.endswith(".keka.com/careers/") else
            _resp(self.KEKA) if "/embedjobs/default/active/b5279857" in url else
            _resp(self.FRESHTEAM) if "freshteam.com/hire/widgets/jobs.json" in url else None)

        def wd_post(url, json):
            self.wd_bodies.append(json)
            if not json["appliedFacets"]:
                return _resp({"total": 11, "jobPostings": [], "facets": self.WD_FACETS})
            return _resp({"total": 2, "jobPostings": self.WD_JOBS[json["offset"]:json["offset"] + 20]})
        self.wd_bodies = []
        client.post.side_effect = wd_post
        return client

    def test_keka(self):
        rows = feeds.fetch_board(JobBoard(ats="keka", token="jupiter", company="Jupiter"), self._client())
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual((r["external_id"], r["location"], r["in_india"], r["level"]),
                         ("jupiter:145340", "Bengaluru, Pune, India", True, "senior"))
        self.assertEqual(r["url"], "https://jupiter.keka.com/careers/jobdetails/145340")
        self.assertIn("AWS", r["description"])

    def test_keka_without_org_id_fails(self):
        with self.assertRaises(ValueError):
            feeds.fetch_board(JobBoard(ats="keka", token="nope", company="Nope"),
                              self._client(keka_page="<title>Invalid Tenant</title>"))

    def test_freshteam(self):
        rows = feeds.fetch_board(JobBoard(ats="freshteam", token="haptik", company="Haptik"), self._client())
        self.assertEqual(len(rows), 1)  # the placeholder "Sample Job Posting" is dropped
        r = rows[0]
        self.assertEqual((r["location"], r["in_india"], r["tags"], r["work_mode"]),
                         ("Mumbai, India", True, ["Engineering"], "onsite"))
        self.assertEqual(r["description"], "Go and Kafka")

    def test_workday_uses_india_facet_and_lazy_description(self):
        board = JobBoard.objects.create(ats="workday", token="nvidia/wd5/NVIDIAExternalCareerSite", company="NVIDIA")
        client = self._client()
        with mock.patch.object(feeds.httpx, "Client") as cls:
            cls.return_value.__enter__.return_value = client
            self.assertEqual(feeds.refresh_boards([board]), (1, 0))
        self.assertEqual(self.wd_bodies[1]["appliedFacets"], {"locationHierarchy1": ["in1"]})
        # No country value: every location naming an Indian city (but not Indianapolis).
        cities = [{"facetParameter": "locationMainGroup", "values": [{"facetParameter": "locations", "values": [
            {"descriptor": "Bangalore, India", "id": "blr"}, {"descriptor": "Indianapolis IN US", "id": "ind"},
            {"descriptor": "Pune", "id": "pnq"}, {"descriptor": "San Jose, CA", "id": "sjc"}]}]}]
        self.assertEqual(feeds._workday_india_facet(cities), ("locations", ["blr", "pnq"]))
        self.assertIsNone(feeds._workday_india_facet([{"facetParameter": "timeType", "values": [
            {"descriptor": "Full time", "id": "ft"}]}]))
        jobs = {j.external_id: j for j in Job.objects.filter(board=board)}
        multi = jobs["nvidia:Research-Scientist_JR2"]
        self.assertEqual((multi.location, multi.in_india), ("India (3 Locations)", True))
        self.assertLess(multi.posted_at, timezone.now() - timedelta(days=29))
        soc = jobs["nvidia:SoC-Architect_JR1"]
        self.assertEqual(soc.url, "https://nvidia.wd5.myworkdayjobs.com/NVIDIAExternalCareerSite/job/India-Bengaluru/SoC-Architect_JR1")
        detail = {"jobPostingInfo": {"jobDescription": "<p>Verilog and Python</p>"}}
        with mock.patch.object(feeds.httpx, "get", return_value=_resp(detail)) as get:
            feeds.ensure_description(soc)
        self.assertEqual(get.call_args.args[0], "https://nvidia.wd5.myworkdayjobs.com/wday/cxs/nvidia/"
                                                "NVIDIAExternalCareerSite/job/India-Bengaluru/SoC-Architect_JR1")
        soc.refresh_from_db()
        self.assertIn("python", soc.skills)

    def test_detect_boards(self):
        page = """
          <a href="https://job-boards.greenhouse.io/razorpaysoftwareprivatelimited/jobs/1">x</a>
          <script src="https://boards.greenhouse.io/embed/job_board/js?for=Druva"></script>
          <a href="https://jobs.lever.co/epifi">x</a> <a href="https://jobs.ashbyhq.com/sarvam">x</a>
          <a href="https://careers.smartrecruiters.com/Freshworks">x</a>
          <a href="https://jupiter.keka.com/careers/">x</a> <link href="https://cdn.keka.com/careers/v/2026/">
          <a href="https://haptik.freshteam.com/jobs">x</a>
          <a href="https://BrowserStack.wd3.myworkdayjobs.com/en-US/External/job/1">x</a>
          <script>{"url":"https:\\/\\/jobs.lever.co\\/dreamsports"}</script>
          <a href="https://wd1.myworkdaysite.com/en-US/recruiting/parexel/Parexel_External_Careers">x</a>
          fetch("https://api.ashbyhq.com/posting-api/job-board/vellum")"""
        self.assertEqual(discovery.detect_boards(page), [
            ("greenhouse", "razorpaysoftwareprivatelimited"), ("greenhouse", "druva"), ("lever", "epifi"),
            ("lever", "dreamsports"), ("ashby", "sarvam"), ("ashby", "vellum"), ("smartrecruiters", "Freshworks"),
            ("keka", "jupiter"), ("freshteam", "haptik"), ("workday", "browserstack/wd3/External"),
            ("workday", "parexel/wd1/Parexel_External_Careers"),
        ])

    def test_slug_guesses(self):
        self.assertEqual(discovery.slug_guesses("Urban Company"), ["urbancompany", "urban-company"])
        self.assertEqual(discovery.slug_guesses("Zoho Technologies Pvt Ltd"), ["zoho"])

    def test_guessed_slug_needs_mostly_india_roles(self):
        rows = [{"in_india": True}, {"in_india": False}, {"in_india": False}]
        with mock.patch.object(feeds, "fetch_board", return_value=rows):
            self.assertIsNone(discovery.verify(discovery.Candidate("lever", "acme", "probe"), "Acme", None))
            found = discovery.verify(discovery.Candidate("lever", "acme", "page"), "Acme", None)
        self.assertEqual((found["jobs"], found["india_jobs"]), (3, 1))
        with mock.patch.object(feeds, "fetch_board", return_value=[{"in_india": False}]):
            self.assertIsNone(discovery.verify(discovery.Candidate("lever", "acme", "page"), "Acme", None))

    def test_any_location_lists_keep_boards_without_india_roles(self):
        with mock.patch.object(feeds, "fetch_board", return_value=[{"in_india": False}]):
            found = discovery.verify(discovery.Candidate("ashby", "acme", "probe"), "Acme", None, india_only=False)
        self.assertEqual((found["jobs"], found["india_jobs"]), (1, 0))
        with mock.patch.object(feeds, "fetch_board", return_value=[]):
            self.assertIsNone(discovery.verify(discovery.Candidate("ashby", "acme", "probe"), "Acme", None, False))

    def test_guessed_keka_board_must_name_the_company(self):
        rows = [{"in_india": True}]
        client = mock.Mock()
        with mock.patch.object(feeds, "fetch_board", return_value=rows):
            client.get.return_value = _resp({"name": "Group Bayport"})
            self.assertIsNone(discovery.verify(discovery.Candidate("keka", "caterpillar", "probe"), "Caterpillar India", client))
            client.get.return_value = _resp({"name": "Workindia"})
            self.assertIsNotNone(discovery.verify(discovery.Candidate("keka", "workindia", "probe"), "WorkIndia", client))

    def test_careers_url_guesses(self):
        self.assertEqual(discovery.careers_url_guesses("Sarvam AI"), [
            "https://sarvamai.com/careers", "https://sarvamai.ai/careers",
            "https://sarvamai.io/careers", "https://sarvamai.in/careers"])
        self.assertNotIn("https://acme.in/careers", discovery.careers_url_guesses("Acme", india_only=False))

    def test_load_companies_from_xlsx(self):
        import tempfile
        from pathlib import Path
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "India Tech Companies"
        ws.append(["No.", "Company", "Careers / search link"])
        ws.append([1, "Infosys", "https://www.google.com/search?q=Infosys+careers"])
        ws.append([2, "Jupiter", "https://jupiter.keka.com/careers/"])
        us = wb.create_sheet("US Startup Targets")
        us.append(["No.", "Company", "Sector"])
        us.append([1, "Acme", "AI"])
        wb.create_sheet("Read Me").append(["Field", "Notes"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "list.xlsx"
            wb.save(path)
            self.assertEqual(discovery.load_companies(path), [
                {"company": "Infosys", "india_only": True},
                {"company": "Jupiter", "india_only": True, "careers_url": "https://jupiter.keka.com/careers/"},
                {"company": "Acme", "india_only": False},
            ])

    def test_companies_with_a_board_are_skipped(self):
        with mock.patch.object(discovery, "known_companies", return_value={"acme"}), \
                mock.patch.object(discovery, "known_boards", return_value=set()), \
                mock.patch.object(discovery, "discover_company") as run:
            run.return_value = discovery.Result("Beta")
            results = discovery.discover([{"company": "Acme Inc"}, {"company": "Beta"}])
        self.assertEqual([(r.company, r.skipped) for r in results], [("Acme Inc", True), ("Beta", False)])
        self.assertEqual(run.call_count, 1)

    def test_add_to_boards_file_inserts_sorted_and_skips_known(self):
        import json
        import tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "boards.json"
            path.write_text('[\n{"ats": "lever", "token": "a", "company": "Alpha"},\n'
                            '{"ats": "lever", "token": "c", "company": "Charlie"}\n]\n')
            with mock.patch.object(feeds, "BOARDS_FILE", path):
                added = discovery.add_to_boards_file([
                    {"ats": "keka", "token": "bravo", "company": "Bravo", "jobs": 1},
                    {"ats": "lever", "token": "A", "company": "Alpha again"},
                ])
            self.assertEqual(added, 1)
            self.assertEqual([e["company"] for e in json.loads(path.read_text())], ["Alpha", "Bravo", "Charlie"])


class CareersAPITests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="u", email="u@ex.com", password="pw12345!")
        self.c = APIClient()
        self.c.force_authenticate(self.user)
        self.job = make_job()
        JobFeedState.objects.create(pk=1, last_refreshed_at=timezone.now())  # skip network

    def test_requires_auth(self):
        self.assertEqual(APIClient().get("/api/careers/jobs/").status_code, 401)

    def test_job_list_and_search(self):
        make_job(title="Chef", description="cook", company="Diner")
        r = self.c.get("/api/careers/jobs/?q=python")
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["ats_quota"], {"limit": 15, "used": 0})

    def test_resume_check_and_daily_cap(self):
        url = f"/api/careers/jobs/{self.job.id}/resume-check/"
        r = self.c.post(url, {"resume_text": RESUME}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertIn("docker", r.data["missing_keywords"])
        for _ in range(14):
            self.assertEqual(self.c.post(url, {"resume_text": RESUME}, format="json").status_code, 200)
        r = self.c.post(url, {"resume_text": RESUME}, format="json")
        self.assertEqual(r.status_code, 429)
        self.assertEqual(r.data["code"], "ats_daily_limit")

    def test_paid_plan_unlimited(self):
        self.user.subscription_plan = "pro"
        self.user.save()
        for i in range(16):
            ResumeCheck.objects.create(user=self.user, job=self.job, ats_score=50)
        r = self.c.post(f"/api/careers/jobs/{self.job.id}/resume-check/", {"resume_text": RESUME}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.data["ats_quota"]["limit"])

    def test_short_resume_rejected_without_consuming_quota(self):
        r = self.c.post(f"/api/careers/jobs/{self.job.id}/resume-check/", {"resume_text": "hi"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(ResumeCheck.objects.count(), 0)

    def test_practice_flow_hides_ideal_answer_until_validated(self):
        Company.objects.all().delete()  # drop the seeded skill bank (migration 0010)
        co = Company.objects.create(name="Acme", is_free=False)
        rnd = Round.objects.create(role=Role.objects.create(company=co, title="SDE"), title="Tech")
        q = InterviewQuestion.objects.create(round=rnd, question_text="Two sum?", question_type="coding",
                                             ideal_answer=PracticeScoringTests.IDEAL)
        lst = self.c.get("/api/careers/practice/questions/")  # accessible: platform is free
        self.assertEqual([x["id"] for x in lst.data], [q.id])
        self.assertNotIn("ideal_answer", lst.data[0])
        r = self.c.post("/api/careers/practice/validate/", {"question_id": q.id, "answer": "hash map"}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertIn("score", r.data)

    @override_settings(PLATFORM_FREE=False)
    def test_practice_respects_plan_when_not_free(self):
        co = Company.objects.create(name="Acme", is_free=False)
        rnd = Round.objects.create(role=Role.objects.create(company=co, title="SDE"), title="Tech")
        InterviewQuestion.objects.create(round=rnd, question_text="Q", question_type="coding", ideal_answer="alpha beta")
        self.assertEqual(self.c.get("/api/careers/practice/questions/").data, [])


@override_settings(PLATFORM_FREE=True)
class PlatformFreeTests(TestCase):
    def test_everyone_gets_max(self):
        from apps.subscriptions.plans import effective_monthly_limit, effective_plan
        for plan in ("free", "pro"):
            u = User.objects.create_user(username=plan, email=f"{plan}@ex.com", password="pw12345!",
                                         subscription_plan=plan)
            self.assertEqual(effective_plan(u), "max")
            self.assertEqual(effective_monthly_limit(u), 50)


class ResumeUploadTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="r", email="r@ex.com", password="pw12345!")
        self.c = APIClient()
        self.c.force_authenticate(self.user)
        self.job = make_job()
        self.url = f"/api/careers/jobs/{self.job.id}/resume-check/"

    def _post(self, name, data):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.c.post(self.url, {"resume_file": SimpleUploadedFile(name, data)}, format="multipart")

    def test_txt_upload(self):
        r = self._post("cv.txt", RESUME.encode())
        self.assertEqual(r.status_code, 200)
        self.assertIn("docker", r.data["missing_keywords"])

    def test_docx_upload(self):
        import io, docx
        d = docx.Document()
        d.add_paragraph(RESUME)
        buf = io.BytesIO()
        d.save(buf)
        r = self._post("cv.docx", buf.getvalue())
        self.assertEqual(r.status_code, 200)
        self.assertIn("python", r.data["matched_keywords"])

    def test_pdf_upload(self):
        from pypdf import PdfWriter
        import io
        w = PdfWriter()
        w.add_blank_page(200, 200)
        buf = io.BytesIO()
        w.write(buf)
        r = self._post("cv.pdf", buf.getvalue())  # blank page -> no text
        self.assertEqual(r.status_code, 400)
        self.assertIn("No readable text", r.data["detail"])

    def test_bad_files_rejected_without_quota(self):
        self.assertEqual(self._post("cv.doc", b"x" * 500).status_code, 400)
        self.assertEqual(self._post("cv.exe", b"MZ" + b"x" * 500).status_code, 400)
        self.assertEqual(self._post("cv.pdf", b"not a pdf at all").status_code, 400)
        self.assertEqual(ResumeCheck.objects.count(), 0)


def _jev_response(fit_score, nouls):
    answers = {"fit": {"type": "score", "score": fit_score}}
    answers.update({f"s{i}": {"type": "noul", "noul": p} for i, p in enumerate(nouls)})
    resp = mock.Mock()
    resp.json.return_value = {"answers": answers}
    resp.raise_for_status.return_value = None
    return resp


class JevTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="j", email="j@ex.com", password="pw12345!")
        self.c = APIClient()
        self.c.force_authenticate(self.user)
        self.job = make_job()  # needs python django postgresql docker aws
        self.url = f"/api/careers/jobs/{self.job.id}/resume-check/"

    @override_settings(TYPESAFE_API_KEY="")
    def test_disabled_without_key_makes_no_call(self):
        with mock.patch("apps.careers.jev.httpx.post") as post:
            r = self.c.post(self.url, {"resume_text": RESUME}, format="json")
        post.assert_not_called()
        self.assertFalse(r.data["ai_assisted"])
        self.assertNotIn("ai_fit", r.data)

    @override_settings(TYPESAFE_API_KEY="k")
    def test_jev_confirms_missing_skill_and_adds_fit(self):
        with override_settings(TYPESAFE_API_KEY=""):
            baseline = self.c.post(self.url, {"resume_text": RESUME}, format="json").data
        missing = baseline["missing_keywords"]  # order matches the Noul questions Jev is asked
        nouls = [0.95 if m == "docker" else 0.1 for m in missing]
        with mock.patch("apps.careers.jev.httpx.post", return_value=_jev_response(3.0, nouls)) as post:
            r = self.c.post(self.url, {"resume_text": RESUME}, format="json")
        self.assertTrue(r.data["ai_assisted"])
        self.assertEqual(r.data["ai_fit"], 75)
        self.assertIn("docker", r.data["matched_keywords"])
        self.assertNotIn("docker", r.data["missing_keywords"])
        self.assertGreater(r.data["ats_score"], baseline["ats_score"])
        self.assertEqual(post.call_args.kwargs["headers"]["Authorization"], "Bearer k")

    @override_settings(TYPESAFE_API_KEY="k")
    def test_jev_failure_falls_back_to_keyword_result(self):
        with mock.patch("apps.careers.jev.httpx.post", side_effect=RuntimeError("boom")):
            r = self.c.post(self.url, {"resume_text": RESUME}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.data["ai_assisted"])
        self.assertIn("docker", r.data["missing_keywords"])

    @override_settings(TYPESAFE_API_KEY="k")
    def test_over_quota_user_never_triggers_jev(self):
        for _ in range(15):
            ResumeCheck.objects.create(user=self.user, job=self.job, ats_score=1)
        with mock.patch("apps.careers.jev.httpx.post") as post:
            r = self.c.post(self.url, {"resume_text": RESUME}, format="json")
        self.assertEqual(r.status_code, 429)
        post.assert_not_called()


class ATSGraphPathTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="g", email="g@ex.com", password="pw12345!")
        self.job = make_job()

    def _run(self, **kw):
        from core.graph import path
        from .graphs import run_ats_check
        kw.setdefault("resume_text", RESUME)
        code, body, trace = run_ats_check(self.user, kw.pop("job_id", self.job.id), **kw)
        return code, body, path(trace)

    @override_settings(TYPESAFE_API_KEY="")
    def test_jev_off_skips_jev_node(self):
        code, _, p = self._run()
        self.assertEqual((code, p), (200, ["parse_resume", "load_job", "keyword_score", "quota_precheck", "save"]))

    @override_settings(TYPESAFE_API_KEY="k")
    def test_jev_on_rescores(self):
        with mock.patch("apps.careers.jev.httpx.post", return_value=_jev_response(4.0, [0.9] * 10)):
            code, body, p = self._run()
        self.assertEqual(p[-3:], ["jev_assess", "rescore", "save"])
        self.assertEqual(body["missing_keywords"], [])

    @override_settings(TYPESAFE_API_KEY="k")
    def test_jev_failure_goes_straight_to_save(self):
        with mock.patch("apps.careers.jev.httpx.post", side_effect=RuntimeError):
            code, body, p = self._run()
        self.assertEqual(p[-2:], ["jev_assess", "save"])
        self.assertFalse(body["ai_assisted"])

    def test_early_exits(self):
        self.assertEqual(self._run(resume_text="short")[::2], (400, ["parse_resume"]))
        self.assertEqual(self._run(job_id=99999)[0], 404)
        nothing = make_job(title="Role", description="Be nice")
        self.assertEqual(self._run(job_id=nothing.id)[0], 422)

    @override_settings(TYPESAFE_API_KEY="k")
    def test_over_quota_ends_before_jev(self):
        for _ in range(15):
            ResumeCheck.objects.create(user=self.user, job=self.job, ats_score=1)
        with mock.patch("apps.careers.jev.httpx.post") as post:
            code, body, p = self._run()
        self.assertEqual((code, p[-1]), (429, "quota_precheck"))
        self.assertIn("ATS Unlimited", body["detail"])
        post.assert_not_called()


PROFILE_RESUME = """Priya Sharma
priya.sharma@example.com | +91 98765 43210 | linkedin.com/in/priyasharma | github.com/priyas
Summary
Backend engineer with 3 years of experience building APIs in Python and Django.
Experience
Software Engineer
Acme Corp
Jan 2023 - Present
- Built REST APIs with Django and PostgreSQL, deployed on AWS with Docker.
Education
B.Tech Computer Science
IIT Delhi
2018 - 2022
Skills
Python, Django, PostgreSQL, Docker, React
"""


class CandidateProfileTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="p", email="p@ex.com", password="pw12345!")
        self.c = APIClient()
        self.c.force_authenticate(self.user)
        JobFeedState.objects.create(pk=1, last_refreshed_at=timezone.now())

    def _upload(self, text=PROFILE_RESUME, name="cv.txt"):
        from django.core.files.uploadedfile import SimpleUploadedFile
        with mock.patch.dict("os.environ", {"OPENROUTER_API_KEY": ""}):
            return self.c.post("/api/careers/profile/resume/",
                               {"resume_file": SimpleUploadedFile(name, text.encode())}, format="multipart")

    def test_upload_fills_profile_with_rules(self):
        r = self._upload()
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["full_name"], "Priya Sharma")
        self.assertEqual(r.data["email"], "priya.sharma@example.com")
        self.assertEqual(r.data["phone"], "+91 98765 43210")
        self.assertIn("linkedin.com/in/priyasharma", r.data["linkedin_url"])
        self.assertEqual(float(r.data["total_experience_years"]), 3.0)
        self.assertTrue({"python", "django", "postgresql", "docker"} <= set(r.data["skills"]))
        exp = r.data["experience"][0]
        self.assertEqual((exp["title"], exp["company"], exp["end"]), ("Software Engineer", "Acme Corp", "Present"))
        self.assertEqual(r.data["education"][0]["institution"], "IIT Delhi")
        self.assertTrue(r.data["has_resume"])
        self.assertEqual(r.data["resume_parsed_by"], "rules")
        dl = self.c.get("/api/careers/profile/resume/")
        self.assertEqual(dl.status_code, 200)
        self.assertIn(b"Priya", b"".join(dl.streaming_content))

    def test_llm_parse_used_when_available(self):
        reply = '```json\n{"full_name": "P S", "skills": ["Go"], "experience": [{"title": "SWE", "company": "X"}], ' \
                '"total_experience_years": 2, "email": "bad-email"}\n```'
        with mock.patch.dict("os.environ", {"OPENROUTER_API_KEY": "k"}), \
                mock.patch("apps.careers.profile_parser.chat_completion", return_value=reply):
            from django.core.files.uploadedfile import SimpleUploadedFile
            r = self.c.post("/api/careers/profile/resume/",
                            {"resume_file": SimpleUploadedFile("cv.txt", PROFILE_RESUME.encode())}, format="multipart")
        self.assertEqual(r.data["resume_parsed_by"], "ai")
        self.assertEqual(r.data["full_name"], "P S")
        self.assertEqual(r.data["skills"][0], "Go")
        self.assertIn("django", r.data["skills"])  # keyword graph tops up the model's list
        self.assertEqual(r.data["email"], "p@ex.com")  # invalid parsed email -> account email

    def test_manual_edit_and_validation(self):
        r = self.c.patch("/api/careers/profile/", {
            "headline": "Backend Engineer", "skills": ["Python", "python", " Go "],
            "experience": [{"title": "SWE", "company": "X", "bogus": "dropped"}],
            "work_modes": ["remote"], "target_roles": ["Backend Engineer"],
        }, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data["skills"], ["Python", "Go"])
        self.assertNotIn("bogus", r.data["experience"][0])
        bad = self.c.patch("/api/careers/profile/", {"work_modes": ["moon"]}, format="json")
        self.assertEqual(bad.status_code, 400)

    def test_bad_resume_rejected(self):
        self.assertEqual(self._upload("too short").status_code, 400)
        self.assertFalse(CandidateProfile.objects.filter(resume_text__gt="").exists())

    def test_ats_check_uses_profile_resume(self):
        self._upload()
        job = make_job()
        r = self.c.post(f"/api/careers/jobs/{job.id}/resume-check/", {}, format="multipart")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertIn("aws", r.data["matched_keywords"])


class SavedAndApplicationsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="a", email="a@ex.com", password="pw12345!")
        self.c = APIClient()
        self.c.force_authenticate(self.user)
        self.job = make_job()
        JobFeedState.objects.create(pk=1, last_refreshed_at=timezone.now())

    def test_save_and_unsave(self):
        self.assertEqual(self.c.post(f"/api/careers/jobs/{self.job.id}/save/").status_code, 200)
        self.c.post(f"/api/careers/jobs/{self.job.id}/save/")  # idempotent
        saved = self.c.get("/api/careers/saved/").data
        self.assertEqual(len(saved), 1)
        self.assertTrue(saved[0]["is_live"])
        self.assertTrue(self.c.get("/api/careers/jobs/").data["results"][0]["saved"])
        self.c.delete(f"/api/careers/jobs/{self.job.id}/save/")
        self.assertEqual(self.c.get("/api/careers/saved/").data, [])

    def test_apply_tracker(self):
        r = self.c.post("/api/careers/applications/", {"job_id": self.job.id}, format="json")
        self.assertEqual(r.status_code, 201)
        self.assertEqual(self.c.post("/api/careers/applications/", {"job_id": self.job.id}, format="json").status_code, 200)
        self.assertEqual(self.c.get("/api/careers/jobs/").data["results"][0]["application_status"], "applied")
        r = self.c.patch(f"/api/careers/applications/{r.data['id']}/", {"status": "interviewing", "notes": "R1 Fri"},
                         format="json")
        self.assertEqual((r.data["status"], r.data["notes"]), ("interviewing", "R1 Fri"))
        manual = self.c.post("/api/careers/applications/", {"title": "SDE", "company": "Elsewhere"}, format="json")
        self.assertEqual(manual.status_code, 201)
        self.assertFalse(manual.data["is_live"])
        self.assertEqual(len(self.c.get("/api/careers/applications/").data), 2)

    def test_other_users_rows_hidden(self):
        app = JobApplication.objects.create(user=User.objects.create_user(username="o", email="o@ex.com"),
                                            title="X", company="Y")
        self.assertEqual(self.c.patch(f"/api/careers/applications/{app.id}/", {"status": "offer"},
                                      format="json").status_code, 404)
        self.assertEqual(self.c.get("/api/careers/applications/").data, [])


class JobFilterAndMatchTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="m", email="m@ex.com", password="pw12345!")
        self.c = APIClient()
        self.c.force_authenticate(self.user)
        JobFeedState.objects.create(pk=1, last_refreshed_at=timezone.now())
        now = timezone.now()
        self.backend = make_job(title="Backend Engineer", skills=["python", "django", "postgresql"],
                                level="mid", work_mode="remote", posted_at=now - timedelta(days=5))
        self.chef = make_job(title="Head Chef", skills=[], level="senior", work_mode="onsite", posted_at=now)
        self.fe = make_job(title="Frontend Engineer", skills=["react", "css"], level="entry", work_mode="hybrid",
                           location="Bengaluru", posted_at=now - timedelta(hours=2))

    def _ids(self, qs):
        return [j["id"] for j in self.c.get("/api/careers/jobs/" + qs).data["results"]]

    def test_filters(self):
        Job.objects.filter(pk=self.fe.pk).update(in_india=True)
        self.assertEqual(self._ids("?region=india"), [self.fe.id])
        self.assertEqual(self._ids("?work_mode=remote"), [self.backend.id])
        self.assertEqual(self._ids("?level=entry,senior"), [self.chef.id, self.fe.id])
        self.assertEqual(self._ids("?posted=1"), [self.chef.id, self.fe.id])
        self.assertEqual(self._ids("?location=bengaluru"), [self.fe.id])
        self.assertEqual(self._ids("?q=django"), [self.backend.id])

    def test_best_match_ranks_by_profile(self):
        r = self.c.get("/api/careers/jobs/?sort=match")
        self.assertEqual(r.data["sort"], "recent")  # no profile yet -> falls back
        self.assertFalse(r.data["has_profile"])
        CandidateProfile.objects.create(user=self.user, skills=["Python", "Django"], target_roles=["Backend Engineer"],
                                        experience_level="mid")
        r = self.c.get("/api/careers/jobs/?sort=match")
        self.assertEqual(r.data["sort"], "match")
        self.assertEqual(r.data["results"][0]["id"], self.backend.id)
        self.assertGreater(r.data["results"][0]["match"], 70)
        self.assertNotIn(self.chef.id, [j["id"] for j in r.data["results"]])  # nothing in common

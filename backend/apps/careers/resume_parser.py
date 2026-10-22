"""Extract plain text from an uploaded resume (PDF, DOCX, or text-like files). Runs locally."""

import io
import re
import zipfile

MAX_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 15
TEXT_EXTS = {"txt", "md", "rtf", "text", "csv", "html", "htm"}


class ResumeParseError(ValueError):
    pass


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ResumeParseError("This PDF is password-protected.")
    return "\n".join((p.extract_text() or "") for p in reader.pages[:MAX_PDF_PAGES])


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for row in t.rows for c in row.cells]
    return "\n".join(parts)


def _plain(data: bytes) -> str:
    text = data.decode("utf-8", errors="ignore")
    text = re.sub(r"<[^>]+>", " ", text)  # html tags
    text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", " ", text) if text.lstrip().startswith("{\\rtf") else text
    return text


def extract_text(filename: str, data: bytes) -> str:
    if len(data) > MAX_BYTES:
        raise ResumeParseError("File is too large (max 5 MB).")
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    try:
        if data[:5] == b"%PDF-" or ext == "pdf":
            text = _pdf(data)
        elif ext == "docx" or (data[:2] == b"PK" and zipfile.is_zipfile(io.BytesIO(data))):
            text = _docx(data)
        elif ext in ("doc", "odt", "pages"):
            raise ResumeParseError("Legacy .doc/.odt files aren't supported — save as PDF or DOCX.")
        elif ext in TEXT_EXTS or ext == "":
            text = _plain(data)
        else:
            raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX or text file.")
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("Couldn't read this file. Try another format or paste the text.") from exc
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 100:
        raise ResumeParseError("No readable text found (scanned image PDFs aren't supported). Paste the text instead.")
    return text
